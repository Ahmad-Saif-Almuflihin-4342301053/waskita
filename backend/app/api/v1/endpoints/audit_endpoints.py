import datetime
import hashlib
import json
import uuid
from typing import List, Optional

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile, status
from sqlalchemy.orm import Session
from starlette.concurrency import run_in_threadpool

from app.core.database import get_db
from app.core.security import get_current_user_optional
from app.models.entities import (
    AuditSession,
    CompetitorDocument,
    CompetitorSnapshot,
    EntityGap,
    MetricResult,
    User,
)
from app.schemas.audit import (
    AnnotatedSpan,
    AuditExecuteRequest,
    AuditExecuteResponse,
    EntityGapItem,
    EntityGapsGroup,
    MetricResultSchema,
    SerpCompetitorItem,
    ThresholdConfig,
)
from app.services.gatekeeper import (
    calculate_composite_cds,
    generate_annotated_spans,
    generate_overview_and_remediation,
)
from app.services.information_gain import calculate_information_gain
from app.services.nlp_pipeline import (
    extract_candidate_entities,
    preprocess_text_async,
    strip_media_placeholders,
)
from app.services.stylometrics import calculate_stylometrics
from app.services.vsm_calculator import compute_vsm_similarity

router = APIRouter(tags=["audit"])


def build_audit_execute_response(session: AuditSession) -> AuditExecuteResponse:
    """Helper untuk memetakan model AuditSession ke AuditExecuteResponse."""
    thresholds_dict = json.loads(session.thresholds_json) if session.thresholds_json else {}
    thresholds = ThresholdConfig(**thresholds_dict) if thresholds_dict else ThresholdConfig()

    annotated_spans = []
    if session.annotated_spans_json:
        raw_spans = json.loads(session.annotated_spans_json)
        annotated_spans = [AnnotatedSpan(**s) for s in raw_spans]

    deficit_items = []
    gain_items = []
    for g in session.entity_gaps:
        item = EntityGapItem(
            term_string=g.term_string,
            term_surface_form=g.term_surface_form,
            tfidf_weight=g.tfidf_weight,
            competitor_consensus_count=g.competitor_consensus_count,
        )
        if g.entity_type == "deficit":
            deficit_items.append(item)
        else:
            gain_items.append(item)

    serp_items = []
    if session.snapshot:
        for doc in session.snapshot.documents:
            serp_items.append(
                SerpCompetitorItem(
                    rank_position=doc.rank_position,
                    page_title=doc.page_title,
                    source_url=doc.source_url,
                    domain_name=doc.domain_name,
                    snippet_text=doc.snippet_text,
                    word_count=doc.word_count,
                    cosine_similarity=doc.cosine_similarity,
                    fetch_status=doc.fetch_status,
                )
            )

    m = session.metrics
    metrics_schema = MetricResultSchema(
        content_differentiation_score=m.content_differentiation_score if m else 0.0,
        avg_cosine_similarity=m.avg_cosine_similarity if m else 0.0,
        max_cosine_similarity=m.max_cosine_similarity if m else 0.0,
        euclidean_distance=m.euclidean_distance if m else 0.0,
        ttr_score=m.ttr_score if m else 0.0,
        mattr_score=m.mattr_score if m else 0.0,
        hko_percentage=m.hko_percentage if m else 0.0,
        is_redundant=m.is_redundant if m else False,
        is_keyword_stuffed=m.is_keyword_stuffed if m else False,
    )

    overview_bullets = [
        f"Audit selesai dengan status vonis {session.verdict_status}.",
        f"Skor diferensiasi semantik CDS tercatat {metrics_schema.content_differentiation_score}%.",
    ]

    return AuditExecuteResponse(
        session_id=session.session_id,
        parent_session_id=session.parent_session_id,
        turn_number=session.turn_number,
        target_query=session.target_query,
        draft_title=session.draft_title,
        draft_content=session.draft_content or "",
        draft_word_count=session.draft_word_count,
        audit_timestamp=session.audit_timestamp.isoformat(),
        verdict_status=session.verdict_status,
        degradation_warning=session.degradation_warning,
        thresholds=thresholds,
        metrics=metrics_schema,
        entity_gaps=EntityGapsGroup(deficit=deficit_items, gain=gain_items),
        annotated_spans=annotated_spans,
        overview_bullets=overview_bullets,
        remediation_prompt=session.remediation_prompt or "",
        serp_competitors=serp_items,
    )


@router.post("/audit/extract-file")
async def extract_draft_from_file(file: UploadFile = File(...)):
    """Ekstraksi teks draf dari berkas yang diunggah (.md, .txt, .docx)."""
    filename = file.filename.lower() if file.filename else ""
    content_bytes = await file.read()

    # Guardrail Ukuran Berkas Maksimal 5 MB (Fase 3 Production Roadmap)
    MAX_FILE_SIZE = 5 * 1024 * 1024  # 5 MB
    if len(content_bytes) > MAX_FILE_SIZE:
        raise HTTPException(
            status_code=status.HTTP_413_REQUEST_ENTITY_TOO_LARGE,
            detail="Ukuran berkas melebihi batas maksimal 5 MB. Mohon unggah dokumen naskah artikel yang lebih ringkas.",
        )

    extracted_text = ""
    if filename.endswith(".txt") or filename.endswith(".md"):
        try:
            extracted_text = content_bytes.decode("utf-8")
        except UnicodeDecodeError:
            extracted_text = content_bytes.decode("latin-1", errors="ignore")
    elif filename.endswith(".docx"):
        try:
            import io
            import docx

            doc = docx.Document(io.BytesIO(content_bytes))
            full_text = []
            for para in doc.paragraphs:
                if para.text.strip():
                    full_text.append(para.text)
            extracted_text = "\n\n".join(full_text)
        except Exception as e:
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail=f"Gagal mengekstrak teks dari berkas DOCX: {str(e)}",
            )
    else:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail="Format berkas tidak didukung. Mohon unggah berkas .md, .txt, atau .docx.",
        )

    words = extracted_text.split()
    file_format = filename.split(".")[-1] if "." in filename else ""
    return {
        "filename": file.filename,
        "draft_content": extracted_text,
        "extracted_text": extracted_text,
        "file_format": file_format,
        "word_count": len(words),
    }


@router.post("/audit/execute", response_model=AuditExecuteResponse)
async def execute_audit(
    req: AuditExecuteRequest,
    current_user: Optional[User] = Depends(get_current_user_optional),
    db: Session = Depends(get_db),
):
    """
    Eksekusi Utama Audit Kualitas Transparan (White-Box IR Engine).
    Mendukung multi-turn audit stream (parent_session_id) dan graceful degradation (FR-13).
    """
    import app.api.v1.routers as routers_contract

    target_query = req.target_query.strip()
    draft_title = req.draft_title.strip()
    draft_content = req.draft_content.strip()
    thresholds = req.thresholds or ThresholdConfig()

    clean_text_for_count = strip_media_placeholders(draft_content)
    words = clean_text_for_count.split()
    draft_word_count = len(words)

    # 1. Tentukan nomor turn jika revisi berulang (Multi-Turn)
    turn_number = 1
    if req.parent_session_id:
        parent = db.query(AuditSession).filter(AuditSession.session_id == req.parent_session_id).first()
        if parent:
            turn_number = parent.turn_number + 1

    # 2. Akuisisi Top 3 SERP Competitors (Scraper / Snapshot Cache / Fallback)
    competitor_docs = None
    if req.use_cached_serp and req.parent_session_id:
        parent_session = (
            db.query(AuditSession)
            .filter(AuditSession.session_id == req.parent_session_id)
            .first()
        )
        if parent_session and parent_session.snapshot and parent_session.snapshot.documents:
            valid_docs = [
                {
                    "rank": d.rank_position,
                    "title": d.page_title,
                    "url": d.source_url,
                    "domain": d.domain_name,
                    "snippet": d.snippet_text or "",
                    "clean_text": d.clean_text or "",
                    "word_count": d.word_count,
                    "fetch_status": d.fetch_status,
                }
                for d in parent_session.snapshot.documents
                if d.clean_text
            ]
            if len(valid_docs) >= 2:
                competitor_docs = valid_docs

    if not competitor_docs:
        competitor_docs = await routers_contract.search_and_scrape_competitors(target_query)

    # FR-13: Graceful Degradation check
    k_docs = len(competitor_docs)
    degradation_warning: Optional[str] = None
    if k_docs < 2:
        raise HTTPException(
            status_code=status.HTTP_503_SERVICE_UNAVAILABLE,
            detail="Korpus kompetitor tidak mencukupi (< 2 halaman dapat diakses dari web atau fallback). "
                   "Mohon coba kembali beberapa saat lagi.",
        )
    elif k_docs == 2:
        degradation_warning = (
            "Perhatian: Hanya 2 dari 3 artikel kompetitor yang berhasil dianalisis. "
            "Kalkulasi VSM dan celah defisit dihitung secara proporsional berdasarkan korpus parsial."
        )

    # 3. CPU Berat #2: Preprocessing & Stemming Korpus via Sastrawi (asyncio.to_thread)
    prep_draft = await preprocess_text_async(draft_content)
    prep_competitors = []
    for c in competitor_docs:
        c_clean = c.get("clean_text", "")
        prep_c = await preprocess_text_async(c_clean)
        prep_competitors.append(prep_c)

    # 4. CPU Berat #3: VSM TF-IDF & Cosine Similarity via run_in_threadpool
    vsm_results = await run_in_threadpool(
        compute_vsm_similarity,
        prep_draft,
        prep_competitors,
    )

    for idx, sim in enumerate(vsm_results["cosine_similarities"]):
        if idx < len(competitor_docs):
            competitor_docs[idx]["cosine_similarity"] = sim

    # 5. Ekstraksi Entitas & Information Gain (Persamaan 2.4 - 2.5)
    draft_entities = extract_candidate_entities(draft_content, top_n=30)
    competitor_entities_list = [
        extract_candidate_entities(c.get("clean_text", ""), top_n=25)
        for c in competitor_docs
    ]
    entity_gaps = calculate_information_gain(
        draft_entities,
        competitor_entities_list,
        draft_tfidf_weights=vsm_results["draft_tfidf_weights"],
    )

    # 6. Analisis Stilometri (TTR, MATTR, HKO)
    stylometrics = calculate_stylometrics(draft_content, target_query)

    # 7. Gatekeeper Aggregator & Rumus Komposit CDS (Bab 2.4)
    draft_stems_set = {e[0] for e in draft_entities}
    serp_stems_set = set()
    for ce in competitor_entities_list:
        for e in ce:
            serp_stems_set.add(e[0])
    shared_entities = draft_stems_set.intersection(serp_stems_set)

    cds_decision = calculate_composite_cds(
        max_cosine_sim=vsm_results["max_cosine_similarity"],
        draft_entities_count=len(draft_stems_set),
        serp_entities_count=len(serp_stems_set),
        shared_entities_count=len(shared_entities),
        gain_count=len(entity_gaps.gain),
        mattr_score=stylometrics["mattr_score"],
        hko_percentage=stylometrics["hko_percentage"],
        thresholds=thresholds,
    )

    # 8. Anotasi Stabilo 4-Warna (Tabel 3.6 Proposal TA)
    annotated_spans = generate_annotated_spans(
        draft_content=draft_content,
        competitor_clean_texts=[c.get("clean_text", "") for c in competitor_docs],
        entity_gaps=entity_gaps,
        hko_percentage=stylometrics["hko_percentage"],
        threshold_hko=thresholds.max_hko_percentage,
    )

    # 9. Poin Editorial & Remediasi Kurasi
    overview_remediation = generate_overview_and_remediation(
        verdict_status=cds_decision["verdict_status"],
        cds_score=cds_decision["content_differentiation_score"],
        is_redundant=cds_decision["is_redundant"],
        is_keyword_stuffed=cds_decision["is_keyword_stuffed"],
        max_cosine_sim=vsm_results["max_cosine_similarity"],
        mattr_score=stylometrics["mattr_score"],
        hko_percentage=stylometrics["hko_percentage"],
        entity_gaps=entity_gaps,
        target_query=target_query,
    )

    # 10. Persistensi ke SQLite WAL (Tabel 3.5 Proposal TA)
    session_id = str(uuid.uuid4())
    now_dt = datetime.datetime.utcnow()

    new_session = AuditSession(
        session_id=session_id,
        user_id=current_user.user_id if current_user else None,
        parent_session_id=req.parent_session_id,
        turn_number=turn_number,
        target_query=target_query,
        draft_title=draft_title,
        draft_content=draft_content,
        draft_word_count=draft_word_count,
        audit_timestamp=now_dt,
        verdict_status=cds_decision["verdict_status"],
        is_pinned=False,
        degradation_warning=degradation_warning,
        annotated_spans_json=json.dumps([s.model_dump() for s in annotated_spans]),
        remediation_prompt=overview_remediation["remediation_prompt"],
        thresholds_json=json.dumps(thresholds.model_dump()),
    )
    db.add(new_session)

    new_metrics = MetricResult(
        session_id=session_id,
        avg_cosine_similarity=vsm_results["avg_cosine_similarity"],
        max_cosine_similarity=vsm_results["max_cosine_similarity"],
        euclidean_distance=vsm_results["euclidean_distance"],
        ttr_score=stylometrics["ttr_score"],
        mattr_score=stylometrics["mattr_score"],
        hko_percentage=stylometrics["hko_percentage"],
        content_differentiation_score=cds_decision["content_differentiation_score"],
        is_redundant=cds_decision["is_redundant"],
        is_keyword_stuffed=cds_decision["is_keyword_stuffed"],
    )
    db.add(new_metrics)

    query_hash = hashlib.sha256(target_query.lower().strip().encode("utf-8")).hexdigest()
    new_snapshot = CompetitorSnapshot(
        session_id=session_id,
        query_hash=query_hash,
        search_engine="Google Indonesia",
        serp_locale="id-ID",
        total_retrieved=len(competitor_docs),
        snapshot_time=now_dt,
    )
    db.add(new_snapshot)
    db.flush()

    serp_competitor_items: List[SerpCompetitorItem] = []
    for c in competitor_docs:
        comp_doc = CompetitorDocument(
            snapshot_id=new_snapshot.snapshot_id,
            rank_position=c.get("rank_position", 1),
            source_url=c.get("source_url", ""),
            domain_name=c.get("domain_name", ""),
            page_title=c.get("page_title", ""),
            snippet_text=c.get("snippet_text", ""),
            clean_text=c.get("clean_text", ""),
            word_count=c.get("word_count", 0),
            cosine_similarity=c.get("cosine_similarity"),
            main_text_hash=c.get("main_text_hash", ""),
            heading_structure=json.dumps(c.get("headings", [])),
            fetch_status=c.get("fetch_status", "success"),
            http_status_code=c.get("http_status_code", 200),
        )
        db.add(comp_doc)

        serp_competitor_items.append(
            SerpCompetitorItem(
                rank_position=c.get("rank_position", 1),
                page_title=c.get("page_title", ""),
                source_url=c.get("source_url", ""),
                domain_name=c.get("domain_name", ""),
                snippet_text=c.get("snippet_text", ""),
                word_count=c.get("word_count", 0),
                cosine_similarity=c.get("cosine_similarity"),
                fetch_status=c.get("fetch_status", "success"),
            )
        )

    for d in entity_gaps.deficit:
        db.add(
            EntityGap(
                session_id=session_id,
                entity_type="deficit",
                term_string=d.term_string,
                term_surface_form=d.term_surface_form,
                tfidf_weight=d.tfidf_weight,
                competitor_consensus_count=d.competitor_consensus_count,
            )
        )
    for g in entity_gaps.gain:
        db.add(
            EntityGap(
                session_id=session_id,
                entity_type="gain",
                term_string=g.term_string,
                term_surface_form=g.term_surface_form,
                tfidf_weight=g.tfidf_weight,
                competitor_consensus_count=g.competitor_consensus_count,
            )
        )

    db.commit()

    return AuditExecuteResponse(
        session_id=session_id,
        parent_session_id=req.parent_session_id,
        turn_number=turn_number,
        target_query=target_query,
        draft_title=draft_title,
        draft_content=draft_content,
        draft_word_count=draft_word_count,
        audit_timestamp=now_dt.isoformat(),
        verdict_status=cds_decision["verdict_status"],
        degradation_warning=degradation_warning,
        thresholds=thresholds,
        metrics=MetricResultSchema(
            content_differentiation_score=cds_decision["content_differentiation_score"],
            avg_cosine_similarity=vsm_results["avg_cosine_similarity"],
            max_cosine_similarity=vsm_results["max_cosine_similarity"],
            euclidean_distance=vsm_results["euclidean_distance"],
            ttr_score=stylometrics["ttr_score"],
            mattr_score=stylometrics["mattr_score"],
            hko_percentage=stylometrics["hko_percentage"],
            is_redundant=cds_decision["is_redundant"],
            is_keyword_stuffed=cds_decision["is_keyword_stuffed"],
        ),
        entity_gaps=entity_gaps,
        annotated_spans=annotated_spans,
        overview_bullets=overview_remediation["overview_bullets"],
        remediation_prompt=overview_remediation["remediation_prompt"],
        serp_competitors=serp_competitor_items,
    )
